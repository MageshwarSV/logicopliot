"""The learned lookup cache: what a kind="lookup" custom field resolved to the one time an
operator typed it in because the material master had nothing for that line, so the same
material/description is filled automatically next time instead of asked for again."""
import io

import openpyxl

from app.core.reference_cache import (
    _fuzzy_similar, lookup_reference, normalize_key, remember_reference,
)
from app.models.custom_field import CustomField
from app.models.custom_field_reference import CustomFieldReferenceValue
from app.models.field_mark import FieldMark
from app.models.job import Job, JobFieldValue
from app.models.template_document import TemplateDocument
from app.models.template_group import TemplateGroup
from tests.conftest import login, make_tenant, make_user


def _make_lookup_field(db_session, tenant, group, label_name="CTH Code", **kwargs):
    cf = CustomField(
        tenant_id=tenant.id, group_id=group.id, label_name=label_name, kind="lookup",
        per_row=True, ask_operator=True, **kwargs,
    )
    db_session.add(cf)
    db_session.commit()
    db_session.refresh(cf)
    return cf


def _make_group(db_session, tenant):
    group = TemplateGroup(tenant_id=tenant.id, name="Sea Import", status="approved")
    db_session.add(group)
    db_session.commit()
    db_session.refresh(group)
    return group


# ---- core logic ------------------------------------------------------------------------------

def test_normalize_key_collapses_spacing_and_case():
    assert normalize_key("  AREC pa  plate  A ") == "AREC PA PLATE A"
    assert normalize_key(None) == ""


def test_lookup_reference_returns_none_when_nothing_learned(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    assert lookup_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"]) is None


def test_remember_then_lookup_round_trips(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"], "85177990")
    db_session.commit()
    assert lookup_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"]) == "85177990"
    # Case/spacing-insensitive, matching material_master's own normalisation.
    assert lookup_reference(db_session, cf.id, ["  839599a", "arec  pa plate a"]) == "85177990"


def test_lookup_reference_is_scoped_to_the_field(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf1 = _make_lookup_field(db_session, tenant, group)
    cf2 = _make_lookup_field(db_session, tenant, group, label_name="RITC Code")
    remember_reference(db_session, cf1.id, ["839599A", "X"], "85177990")
    db_session.commit()
    assert lookup_reference(db_session, cf2.id, ["839599A", "X"]) is None


def test_remember_reference_ignores_a_blank_answer(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "   ")
    db_session.commit()
    assert db_session.query(CustomFieldReferenceValue).count() == 0


def test_remember_reference_ignores_values_with_nothing_to_key_on(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["", ""], "85177990")
    db_session.commit()
    assert db_session.query(CustomFieldReferenceValue).count() == 0


def test_remember_reference_updates_an_existing_row_instead_of_duplicating(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177990")
    db_session.commit()
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177991")
    db_session.commit()
    assert db_session.query(CustomFieldReferenceValue).count() == 1
    assert lookup_reference(db_session, cf.id, ["839599A", "X"]) == "85177991"


# ---- learning through a correction ------------------------------------------------------------

def test_correcting_a_blank_lookup_value_is_learned(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    job = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-LEARN1", status="extracted")
    db_session.add(job)
    db_session.commit()
    db_session.add_all([
        JobFieldValue(tenant_id=tenant.id, job_id=job.id, label_name="item_material_code",
                     extracted_value="839599A", set_index=1, row_index=1),
        JobFieldValue(tenant_id=tenant.id, job_id=job.id, label_name="product_description",
                     extracted_value="AREC PA PLATE A", set_index=1, row_index=1),
    ])
    fv = JobFieldValue(tenant_id=tenant.id, job_id=job.id, custom_field_id=cf.id,
                       label_name="CTH Code", extracted_value="", set_index=1, row_index=1)
    db_session.add(fv)
    db_session.commit()
    db_session.refresh(fv)

    op = make_user(db_session, role="operator", tenant=tenant, email="op-learn1@example.com")
    login(client, op.email)

    resp = client.patch(f"/api/v1/job-field-values/{fv.id}", json={"corrected_value": "85177990"})
    assert resp.status_code == 200, resp.text

    assert lookup_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"]) == "85177990"


def test_correcting_a_value_the_master_already_found_is_not_learned(client, db_session):
    """The dump data already answered this line - a one-off manual override of it must not
    poison every future line that legitimately matches the dump."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    job = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-LEARN2", status="extracted")
    db_session.add(job)
    db_session.commit()
    fv = JobFieldValue(tenant_id=tenant.id, job_id=job.id, custom_field_id=cf.id,
                       label_name="CTH Code", extracted_value="85177990",
                       set_index=1, row_index=1)
    db_session.add(fv)
    db_session.commit()
    db_session.refresh(fv)

    op = make_user(db_session, role="operator", tenant=tenant, email="op-learn2@example.com")
    login(client, op.email)

    resp = client.patch(f"/api/v1/job-field-values/{fv.id}", json={"corrected_value": "99999999"})
    assert resp.status_code == 200, resp.text
    assert db_session.query(CustomFieldReferenceValue).count() == 0


def test_correcting_a_non_lookup_custom_field_is_not_learned(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = CustomField(tenant_id=tenant.id, group_id=group.id, label_name="Note", kind="ai")
    db_session.add(cf)
    db_session.commit()
    job = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-LEARN3", status="extracted")
    db_session.add(job)
    db_session.commit()
    fv = JobFieldValue(tenant_id=tenant.id, job_id=job.id, custom_field_id=cf.id,
                       label_name="Note", extracted_value="")
    db_session.add(fv)
    db_session.commit()
    db_session.refresh(fv)

    op = make_user(db_session, role="operator", tenant=tenant, email="op-learn3@example.com")
    login(client, op.email)

    resp = client.patch(f"/api/v1/job-field-values/{fv.id}", json={"corrected_value": "hello"})
    assert resp.status_code == 200, resp.text
    assert db_session.query(CustomFieldReferenceValue).count() == 0


# ---- admin endpoints ---------------------------------------------------------------------------

def test_list_reference_values(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177990")
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-list@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/custom-fields/{cf.id}/reference-values")
    assert resp.status_code == 200, resp.text
    values = resp.json()["values"]
    assert len(values) == 1
    assert values[0]["match_value_1"] == "839599A"
    assert values[0]["resolved_value"] == "85177990"


def test_list_reference_values_rejects_a_non_lookup_field(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = CustomField(tenant_id=tenant.id, group_id=group.id, label_name="Note", kind="ai")
    db_session.add(cf)
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-reject@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/custom-fields/{cf.id}/reference-values")
    assert resp.status_code == 400


def test_download_template_has_the_right_columns(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)

    admin = make_user(db_session, role="super_admin", email="sa-template@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/custom-fields/{cf.id}/reference-values/template")
    assert resp.status_code == 200, resp.text
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header == ["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4",
                       "Resolved Value"]


def test_upload_reference_values_bulk_loads_rows(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)

    admin = make_user(db_session, role="super_admin", email="sa-upload@example.com")
    login(client, admin.email)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4", "Resolved Value"])
    ws.append(["839599A", "AREC PA PLATE A", "", "", "85177990"])
    ws.append(["094624A", "", "", "", "85340000"])
    ws.append(["", "", "", "", ""])  # blank row, skipped
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.post(
        f"/api/v1/custom-fields/{cf.id}/reference-values/upload",
        files={"file": ("template.xlsx", buffer.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["rows_added"] == 2
    assert body["rows_skipped"] == 1
    assert body["conflicts"] == []
    assert lookup_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"]) == "85177990"
    assert lookup_reference(db_session, cf.id, ["094624A", ""]) == "85340000"


def test_upload_flags_a_conflicting_row_without_overwriting_it(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"], "85177990")
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-conflict@example.com")
    login(client, admin.email)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4", "Resolved Value"])
    ws.append(["839599A", "AREC PA PLATE A", "", "", "99999999"])  # different from what is on file
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.post(
        f"/api/v1/custom-fields/{cf.id}/reference-values/upload",
        files={"file": ("template.xlsx", buffer.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["rows_added"] == 0
    assert len(body["conflicts"]) == 1
    conflict = body["conflicts"][0]
    assert conflict["existing_value"] == "85177990"
    assert conflict["uploaded_value"] == "99999999"
    # Untouched - the upload never applies its own value over an existing different one.
    assert lookup_reference(db_session, cf.id, ["839599A", "AREC PA PLATE A"]) == "85177990"


def test_upload_of_an_identical_value_is_unchanged_not_a_conflict(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177990")
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-samevalue@example.com")
    login(client, admin.email)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4", "Resolved Value"])
    ws.append(["839599A", "X", "", "", "85177990"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.post(
        f"/api/v1/custom-fields/{cf.id}/reference-values/upload",
        files={"file": ("template.xlsx", buffer.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["rows_unchanged"] == 1
    assert body["conflicts"] == []


def test_patch_reference_value_applies_the_uploaded_value(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177990")
    db_session.commit()
    row_id = db_session.query(CustomFieldReferenceValue).one().id

    admin = make_user(db_session, role="super_admin", email="sa-patch@example.com")
    login(client, admin.email)

    resp = client.patch(f"/api/v1/custom-fields/{cf.id}/reference-values/{row_id}",
                        json={"resolved_value": "99999999"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["resolved_value"] == "99999999"
    assert lookup_reference(db_session, cf.id, ["839599A", "X"]) == "99999999"


def test_upload_reference_values_rejects_a_non_excel_file(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)

    admin = make_user(db_session, role="super_admin", email="sa-reject2@example.com")
    login(client, admin.email)

    resp = client.post(
        f"/api/v1/custom-fields/{cf.id}/reference-values/upload",
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 422


def test_delete_reference_value(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["839599A", "X"], "85177990")
    db_session.commit()
    row_id = db_session.query(CustomFieldReferenceValue).one().id

    admin = make_user(db_session, role="super_admin", email="sa-delete@example.com")
    login(client, admin.email)

    resp = client.delete(f"/api/v1/custom-fields/{cf.id}/reference-values/{row_id}")
    assert resp.status_code == 204, resp.text
    assert db_session.query(CustomFieldReferenceValue).count() == 0


def test_list_reference_values_accepts_a_target_value_custom_field(client, db_session):
    """kind="ai" is not "lookup", but is_target_value=True must still be let through - the
    gate was broadened for exactly this case (Target Value feature, Phase 2)."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = CustomField(tenant_id=tenant.id, group_id=group.id, label_name="Quotation",
                     kind="ai", is_target_value=True)
    db_session.add(cf)
    db_session.commit()
    remember_reference(db_session, cf.id, ["RFQ/0003/23-24"], "RFQ/0003/23-24")
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-tv-cf@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/custom-fields/{cf.id}/reference-values")
    assert resp.status_code == 200, resp.text
    assert resp.json()["values"][0]["resolved_value"] == "RFQ/0003/23-24"


# ---- mark-scoped reference values (Target Value feature, Phase 2) -----------------------------

def _make_doc(db_session, tenant, group):
    doc = TemplateDocument(tenant_id=tenant.id, group_id=group.id, name="Invoice", doc_type="Invoice")
    db_session.add(doc)
    db_session.commit()
    db_session.refresh(doc)
    return doc


def _make_target_value_mark(db_session, tenant, doc, label_name="IEC Number"):
    mark = FieldMark(tenant_id=tenant.id, document_id=doc.id, label_name=label_name,
                     page_number=1, x=0.1, y=0.1, width=0.1, height=0.1, is_target_value=True)
    db_session.add(mark)
    db_session.commit()
    db_session.refresh(mark)
    return mark


def test_list_mark_reference_values(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    remember_reference(db_session, None, ["0301234567"], "0301234567 (Verified)", mark_id=mark.id)
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-mark-list@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/marks/{mark.id}/reference-values")
    assert resp.status_code == 200, resp.text
    values = resp.json()["values"]
    assert len(values) == 1
    assert values[0]["match_value_1"] == "0301234567"
    assert values[0]["resolved_value"] == "0301234567 (Verified)"


def test_list_mark_reference_values_rejects_a_non_target_value_mark(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = FieldMark(tenant_id=tenant.id, document_id=doc.id, label_name="Plain field",
                     page_number=1, x=0.1, y=0.1, width=0.1, height=0.1)
    db_session.add(mark)
    db_session.commit()
    db_session.refresh(mark)

    admin = make_user(db_session, role="super_admin", email="sa-mark-reject@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/marks/{mark.id}/reference-values")
    assert resp.status_code == 400


def test_mark_and_custom_field_reference_values_do_not_leak_into_each_other(client, db_session):
    """The generalised owner filter (mark_id vs custom_field_id) must keep a mark's table and a
    custom field's table completely separate, even when both happen to share a match value."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, None, ["SHARED-KEY"], "from mark table", mark_id=mark.id)
    remember_reference(db_session, cf.id, ["SHARED-KEY"], "from custom field table")
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-no-leak@example.com")
    login(client, admin.email)

    mark_values = client.get(f"/api/v1/marks/{mark.id}/reference-values").json()["values"]
    cf_values = client.get(f"/api/v1/custom-fields/{cf.id}/reference-values").json()["values"]
    assert len(mark_values) == 1 and mark_values[0]["resolved_value"] == "from mark table"
    assert len(cf_values) == 1 and cf_values[0]["resolved_value"] == "from custom field table"
    assert lookup_reference(db_session, None, ["SHARED-KEY"], mark_id=mark.id) == "from mark table"
    assert lookup_reference(db_session, cf.id, ["SHARED-KEY"]) == "from custom field table"


def test_download_mark_template_has_the_right_columns(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)

    admin = make_user(db_session, role="super_admin", email="sa-mark-template@example.com")
    login(client, admin.email)

    resp = client.get(f"/api/v1/marks/{mark.id}/reference-values/template")
    assert resp.status_code == 200, resp.text
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header == ["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4",
                       "Resolved Value"]


def test_upload_mark_reference_values_bulk_loads_rows(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)

    admin = make_user(db_session, role="super_admin", email="sa-mark-upload@example.com")
    login(client, admin.email)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4", "Resolved Value"])
    ws.append(["0301234567", "", "", "", "0301234567 (Verified)"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.post(
        f"/api/v1/marks/{mark.id}/reference-values/upload",
        files={"file": ("template.xlsx", buffer.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["rows_added"] == 1
    assert body["conflicts"] == []
    assert lookup_reference(db_session, None, ["0301234567"], mark_id=mark.id) == "0301234567 (Verified)"


def test_upload_mark_reference_values_flags_a_conflict(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    remember_reference(db_session, None, ["0301234567"], "OLD VALUE", mark_id=mark.id)
    db_session.commit()

    admin = make_user(db_session, role="super_admin", email="sa-mark-conflict@example.com")
    login(client, admin.email)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Match Value 1", "Match Value 2", "Match Value 3", "Match Value 4", "Resolved Value"])
    ws.append(["0301234567", "", "", "", "NEW VALUE"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.post(
        f"/api/v1/marks/{mark.id}/reference-values/upload",
        files={"file": ("template.xlsx", buffer.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["rows_added"] == 0
    assert len(body["conflicts"]) == 1
    assert body["conflicts"][0]["existing_value"] == "OLD VALUE"
    assert body["conflicts"][0]["uploaded_value"] == "NEW VALUE"
    assert lookup_reference(db_session, None, ["0301234567"], mark_id=mark.id) == "OLD VALUE"


def test_patch_mark_reference_value_applies_the_uploaded_value(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    remember_reference(db_session, None, ["0301234567"], "OLD VALUE", mark_id=mark.id)
    db_session.commit()
    row_id = db_session.query(CustomFieldReferenceValue).filter_by(mark_id=mark.id).one().id

    admin = make_user(db_session, role="super_admin", email="sa-mark-patch@example.com")
    login(client, admin.email)

    resp = client.patch(f"/api/v1/marks/{mark.id}/reference-values/{row_id}",
                        json={"resolved_value": "NEW VALUE"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["resolved_value"] == "NEW VALUE"
    assert lookup_reference(db_session, None, ["0301234567"], mark_id=mark.id) == "NEW VALUE"


def test_patch_mark_reference_value_via_the_custom_field_url_is_rejected(client, db_session):
    """The owner filter must be enforced on the URL used, not just on which row id is passed -
    a mark's row id must not be reachable/editable through the custom-fields endpoint."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, None, ["0301234567"], "OLD VALUE", mark_id=mark.id)
    db_session.commit()
    row_id = db_session.query(CustomFieldReferenceValue).filter_by(mark_id=mark.id).one().id

    admin = make_user(db_session, role="super_admin", email="sa-mark-crossurl@example.com")
    login(client, admin.email)

    resp = client.patch(f"/api/v1/custom-fields/{cf.id}/reference-values/{row_id}",
                        json={"resolved_value": "HIJACKED"})
    assert resp.status_code == 404


def test_delete_mark_reference_value(client, db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    doc = _make_doc(db_session, tenant, group)
    mark = _make_target_value_mark(db_session, tenant, doc)
    remember_reference(db_session, None, ["0301234567"], "OLD VALUE", mark_id=mark.id)
    db_session.commit()
    row_id = db_session.query(CustomFieldReferenceValue).filter_by(mark_id=mark.id).one().id

    admin = make_user(db_session, role="super_admin", email="sa-mark-delete@example.com")
    login(client, admin.email)

    resp = client.delete(f"/api/v1/marks/{mark.id}/reference-values/{row_id}")
    assert resp.status_code == 204, resp.text
    assert db_session.query(CustomFieldReferenceValue).count() == 0


# ---- fuzzy matching (opt-in, is_target_value fields only) -------------------------------------

def test_fuzzy_similar_exact_match():
    assert _fuzzy_similar("KUEHNE+NAGEL", "KUEHNE+NAGEL") is True


def test_fuzzy_similar_ignores_spacing_and_punctuation():
    assert _fuzzy_similar("KUEHNE + NAGEL", "KUEHNE+NAGEL") is True
    assert _fuzzy_similar("Kuehne & Nagel", "KUEHNE NAGEL") is True


def test_fuzzy_similar_matches_abbreviation_against_full_legal_name():
    # The most common real case: the same company printed with and without its legal suffix.
    assert _fuzzy_similar("KUEHNE+NAGEL", "KUEHNE + NAGEL PVT. LTD.") is True
    assert _fuzzy_similar("KUEHNE + NAGEL PVT. LTD.", "KUEHNE+NAGEL CO., LTD.") is True


def test_fuzzy_similar_tolerates_ocr_noise():
    # One misread character in an otherwise-identical name - close enough to pass the
    # SequenceMatcher fallback without qualifying for the substring shortcut.
    assert _fuzzy_similar("EASTERN LINER SHIPPING PRIVATE LIMITED",
                          "EASTERN LINER SH1PPING PRIVATE LIMITED") is True


def test_fuzzy_similar_rejects_a_genuinely_different_company():
    assert _fuzzy_similar("KUEHNE + NAGEL PVT. LTD.", "EVERGREEN SHIPPING AGENCY") is False
    assert _fuzzy_similar("", "KUEHNE+NAGEL") is False


def test_lookup_reference_default_stays_exact_even_for_a_near_miss(db_session):
    """The default (fuzzy=False) must never soften - this is the same lookup_reference a
    kind="lookup" CTH field calls, and that one must never fuzzy-match."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["KUEHNE + NAGEL PVT. LTD."], "RFQ/0003/23-24")
    db_session.commit()

    assert lookup_reference(db_session, cf.id, ["KUEHNE+NAGEL"]) is None
    assert lookup_reference(db_session, cf.id, ["KUEHNE+NAGEL"], fuzzy=False) is None


def test_lookup_reference_fuzzy_matches_a_differently_spelled_key(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["KUEHNE + NAGEL PVT. LTD."], "RFQ/0003/23-24")
    db_session.commit()

    assert lookup_reference(db_session, cf.id, ["KUEHNE+NAGEL"], fuzzy=True) == "RFQ/0003/23-24"
    assert lookup_reference(db_session, cf.id, ["Kuehne & Nagel"], fuzzy=True) == "RFQ/0003/23-24"


def test_lookup_reference_fuzzy_does_not_match_an_unrelated_company(db_session):
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = _make_lookup_field(db_session, tenant, group)
    remember_reference(db_session, cf.id, ["KUEHNE + NAGEL PVT. LTD."], "RFQ/0003/23-24")
    db_session.commit()

    assert lookup_reference(db_session, cf.id, ["EVERGREEN SHIPPING AGENCY"], fuzzy=True) is None


def test_target_value_field_with_fuzzy_match_resolves_a_spelling_variant_end_to_end(client, db_session):
    """The real scenario this was built for: an operator answers Quotation Value once for
    'KUEHNE + NAGEL PVT. LTD.', and a LATER job whose document happens to print the company's
    name as plain 'KUEHNE+NAGEL' still auto-resolves instead of asking again."""
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    cf = CustomField(tenant_id=tenant.id, group_id=group.id, label_name="Quotation Value",
                     kind="ai", is_target_value=True, fuzzy_match=True,
                     ask_operator=True, ask_operator_required=True)
    db_session.add(cf)
    db_session.commit()
    db_session.refresh(cf)

    job1 = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-FUZZY1", status="extracted")
    db_session.add(job1)
    db_session.commit()
    fv1 = JobFieldValue(tenant_id=tenant.id, job_id=job1.id, custom_field_id=cf.id,
                        label_name="Quotation Value", extracted_value="",
                        target_value_raw="KUEHNE + NAGEL PVT. LTD.")
    db_session.add(fv1)
    db_session.commit()
    db_session.refresh(fv1)

    op = make_user(db_session, role="operator", tenant=tenant, email="op-fuzzy@example.com")
    login(client, op.email)
    resp = client.patch(f"/api/v1/job-field-values/{fv1.id}", json={"corrected_value": "RFQ/0003/23-24"})
    assert resp.status_code == 200, resp.text

    # A later job's raw extraction differs only in spacing/punctuation from what was just taught.
    assert lookup_reference(db_session, custom_field_id=cf.id, match_values=["KUEHNE+NAGEL"],
                            fuzzy=True) == "RFQ/0003/23-24"
    # And confirms it would NOT have resolved without fuzzy=True.
    assert lookup_reference(db_session, custom_field_id=cf.id, match_values=["KUEHNE+NAGEL"],
                            fuzzy=False) is None
