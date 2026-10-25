"""GET /finance/invoices (cross-tenant search) and /finance/reports/<type>[/export] - row
shapes and that the exported files are actually parseable, not just that a 200 came back."""
import csv
import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from app.models.billable_page_event import BillablePageEvent
from app.models.invoice import Invoice
from app.models.job import Job, JobDocument
from app.models.template_document import TemplateDocument
from app.models.template_group import TemplateGroup
from app.models.tenant_billing_config import TenantBillingConfig
from tests.conftest import login, make_tenant, make_user


def _make_group(db_session, tenant, mode="Sea Import"):
    group = TemplateGroup(tenant_id=tenant.id, name="Sea Import", status="approved", mode=mode)
    db_session.add(group)
    db_session.commit()
    db_session.refresh(group)
    return group


def _make_template_document(db_session, tenant, group, name="BL"):
    doc = TemplateDocument(tenant_id=tenant.id, group_id=group.id, name=name, doc_type="BL")
    db_session.add(doc)
    db_session.commit()
    db_session.refresh(doc)
    return doc


def _make_job_with_pages(db_session, tenant, group, template_doc, *, pages, created_at):
    job = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-FINRPT", status="completed")
    db_session.add(job)
    db_session.commit()
    db_session.refresh(job)
    job.created_at = created_at
    doc = JobDocument(
        tenant_id=tenant.id, job_id=job.id, template_document_id=template_doc.id,
        file_path="/tmp/fake.pdf", page_count=pages,
    )
    db_session.add(doc)
    if pages:
        event = BillablePageEvent(
            tenant_id=tenant.id, job_id=job.id, job_reference=job.reference,
            shipment_type=group.mode or "Unassigned", page_count=pages,
        )
        db_session.add(event)
        db_session.commit()
        event.created_at = created_at
    db_session.commit()
    return job


def _save_billing_config(db_session, tenant, *, created_at=None, **overrides):
    config = TenantBillingConfig(tenant_id=tenant.id, **overrides)
    db_session.add(config)
    db_session.commit()
    if created_at is not None:
        config.created_at = created_at
        db_session.commit()
    db_session.refresh(config)
    return config


def _setup_one_finalized_invoice(db_session):
    now = datetime.now(timezone.utc)
    last_month_day = now.replace(day=1) - timedelta(days=1)
    tenant = make_tenant(db_session)
    group = _make_group(db_session, tenant)
    tdoc = _make_template_document(db_session, tenant, group)
    _make_job_with_pages(db_session, tenant, group, tdoc, pages=40, created_at=last_month_day)
    _save_billing_config(db_session, tenant, created_at=last_month_day.replace(day=1))
    return tenant


def test_list_finance_invoices_across_tenants(client, db_session):
    tenant = _setup_one_finalized_invoice(db_session)
    fin = make_user(db_session, role="saas_finance", email="fin-rpt1@example.com")
    login(client, fin.email)

    resp = client.get("/api/v1/finance/invoices")
    assert resp.status_code == 200, resp.text
    rows = resp.json()
    assert len(rows) == 1
    assert rows[0]["tenant_id"] == tenant.id
    assert rows[0]["tenant_name"] == tenant.name
    assert rows[0]["status"] == "unpaid"


def test_list_finance_invoices_filters_by_status(client, db_session):
    tenant = _setup_one_finalized_invoice(db_session)
    fin = make_user(db_session, role="saas_finance", email="fin-rpt2@example.com")
    login(client, fin.email)
    client.get("/api/v1/finance/invoices")  # ensure finalized
    invoice = db_session.query(Invoice).filter(Invoice.tenant_id == tenant.id).first()
    client.post(f"/api/v1/tenants/{tenant.id}/billing/invoices/{invoice.id}/mark-paid", json={"paid": True})

    unpaid_resp = client.get("/api/v1/finance/invoices?status=unpaid")
    paid_resp = client.get("/api/v1/finance/invoices?status=paid")
    assert unpaid_resp.json() == []
    assert len(paid_resp.json()) == 1


def test_customer_billing_report_rows(client, db_session):
    _setup_one_finalized_invoice(db_session)
    fin = make_user(db_session, role="saas_finance", email="fin-rpt3@example.com")
    login(client, fin.email)
    resp = client.get("/api/v1/finance/reports/customer-billing")
    assert resp.status_code == 200, resp.text
    rows = resp.json()
    assert len(rows) == 1
    assert "Platform Fee" in rows[0]
    assert "Total Amount" in rows[0]


def test_unknown_report_type_404s(client, db_session):
    fin = make_user(db_session, role="saas_finance", email="fin-rpt4@example.com")
    login(client, fin.email)
    resp = client.get("/api/v1/finance/reports/not-a-real-report")
    assert resp.status_code == 404


def test_csv_export_is_parseable(client, db_session):
    _setup_one_finalized_invoice(db_session)
    fin = make_user(db_session, role="saas_finance", email="fin-rpt5@example.com")
    login(client, fin.email)
    resp = client.get("/api/v1/finance/reports/customer-billing/export?format=csv")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    text = resp.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    assert rows[0] == ["Customer", "Jobs", "Pages", "Platform Fee", "Usage Charges", "GST", "Total Amount", "Invoice Status"]
    assert len(rows) == 2  # header + one tenant


def test_xlsx_export_round_trips(client, db_session):
    _setup_one_finalized_invoice(db_session)
    fin = make_user(db_session, role="saas_finance", email="fin-rpt6@example.com")
    login(client, fin.email)
    resp = client.get("/api/v1/finance/reports/customer-billing/export?format=xlsx")
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Customer", "Jobs", "Pages", "Platform Fee", "Usage Charges", "GST", "Total Amount", "Invoice Status")
    assert len(rows) == 2


def test_tenant_admin_cannot_export_reports(client, db_session):
    tenant = make_tenant(db_session)
    ta = make_user(db_session, role="tenant_admin", tenant=tenant, email="ta-rpt1@example.com")
    login(client, ta.email)
    resp = client.get("/api/v1/finance/reports/customer-billing/export?format=csv")
    assert resp.status_code == 403
