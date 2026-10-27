"""upload_job_document (the manual per-slot upload, as opposed to Smart Upload) now accepts
Word/Excel files too, converting them to PDF before landing them in the chosen slot - found
live: an operator manually placing a document into a specific slot (rather than using Smart
Upload's auto-classify) had no way to do that for a Word/Excel file at all, even though Smart
Upload already converted the exact same file types. Mirrors Smart Upload's own office_convert
usage, scoped to exactly one output file per call since this endpoint always lands ONE file
into ONE chosen slot - a multi-sheet Excel workbook (which would become several PDFs) has
nowhere single to go and is rejected with a message pointing at Smart Upload instead."""
import subprocess
from unittest.mock import patch

import fitz
import openpyxl
from app.models.field_mark import FieldMark
from app.models.job import Job, JobDocument
from app.models.template_document import TemplateDocument
from app.models.template_group import TemplateGroup
from tests.conftest import login, make_tenant, make_user


def _make_template(db_session, tenant):
    group = TemplateGroup(tenant_id=tenant.id, name="Sea Import", status="approved")
    db_session.add(group)
    db_session.commit()
    db_session.refresh(group)
    tdoc = TemplateDocument(tenant_id=tenant.id, group_id=group.id, name="Invoice", doc_type="invoice")
    db_session.add(tdoc)
    db_session.commit()
    db_session.refresh(tdoc)
    mark = FieldMark(tenant_id=tenant.id, document_id=tdoc.id, label_name="Invoice No",
                     page_number=1, x=0.1, y=0.1, width=0.2, height=0.05)
    db_session.add(mark)
    db_session.commit()
    return group, tdoc


def _make_job(db_session, tenant, group, tdoc):
    job = Job(tenant_id=tenant.id, group_id=group.id, reference="JOB-OFFICE", status="draft")
    db_session.add(job)
    db_session.commit()
    db_session.refresh(job)
    slot = JobDocument(tenant_id=tenant.id, job_id=job.id, template_document_id=tdoc.id,
                       file_path=None, page_count=0, file_index=0)
    db_session.add(slot)
    db_session.commit()
    return job


def _login_operator(client, db_session, tenant):
    op = make_user(db_session, role="operator", tenant=tenant, email="op-office@example.com")
    login(client, op.email)
    return op


def _fake_soffice_run(args, **kwargs):
    """Stands in for the real LibreOffice call - writes a genuinely openable tiny PDF at the
    exact path _run_soffice's own --outdir/src args say it should land at, same technique
    test_office_convert_timeout_retry.py uses for the lower-level function directly."""
    from pathlib import Path

    out_dir = Path(args[args.index("--outdir") + 1])
    src = Path(args[-1])
    doc = fitz.open()
    doc.new_page(width=200, height=200)
    doc.save(out_dir / f"{src.stem}.pdf")
    return subprocess.CompletedProcess(args, 0, stdout=b"", stderr=b"")


def _xlsx_bytes(sheet_names: list[str]) -> bytes:
    import io

    wb = openpyxl.Workbook()
    wb.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_word_file_is_converted_to_pdf_and_lands_in_the_slot(client, db_session):
    tenant = make_tenant(db_session)
    group, tdoc = _make_template(db_session, tenant)
    job = _make_job(db_session, tenant, group, tdoc)
    _login_operator(client, db_session, tenant)

    with patch("app.core.office_convert.subprocess.run", side_effect=_fake_soffice_run):
        resp = client.post(
            f"/api/v1/jobs/{job.id}/documents/{tdoc.id}/upload",
            files={"file": ("invoice.docx", b"not a real docx - soffice is mocked",
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 200, resp.text
    doc = next(d for d in resp.json()["documents"] if d["id"] == tdoc.id or d.get("name") == "Invoice")
    assert doc["is_uploaded"] is True
    # Converted content is landed as a PDF, not the original Word file.
    assert doc["original_name"].endswith(".pdf")


def test_a_single_sheet_excel_file_is_converted_to_pdf_and_lands_in_the_slot(client, db_session):
    tenant = make_tenant(db_session)
    group, tdoc = _make_template(db_session, tenant)
    job = _make_job(db_session, tenant, group, tdoc)
    _login_operator(client, db_session, tenant)

    with patch("app.core.office_convert.subprocess.run", side_effect=_fake_soffice_run):
        resp = client.post(
            f"/api/v1/jobs/{job.id}/documents/{tdoc.id}/upload",
            files={"file": ("invoice.xlsx", _xlsx_bytes(["Sheet1"]),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 200, resp.text
    doc = next(d for d in resp.json()["documents"] if d.get("name") == "Invoice")
    assert doc["is_uploaded"] is True
    assert doc["original_name"].endswith(".pdf")


def test_a_multi_sheet_excel_file_is_rejected_with_a_clear_message(client, db_session):
    tenant = make_tenant(db_session)
    group, tdoc = _make_template(db_session, tenant)
    job = _make_job(db_session, tenant, group, tdoc)
    _login_operator(client, db_session, tenant)

    with patch("app.core.office_convert.subprocess.run", side_effect=_fake_soffice_run):
        resp = client.post(
            f"/api/v1/jobs/{job.id}/documents/{tdoc.id}/upload",
            files={"file": ("invoice.xlsx", _xlsx_bytes(["Sheet1", "Sheet2"]),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 422, resp.text
    assert "2 sheets" in resp.json()["detail"]
    assert "Smart Upload" in resp.json()["detail"]


def test_a_conversion_failure_surfaces_the_office_convert_error_message(client, db_session):
    tenant = make_tenant(db_session)
    group, tdoc = _make_template(db_session, tenant)
    job = _make_job(db_session, tenant, group, tdoc)
    _login_operator(client, db_session, tenant)

    def fake_run(args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="soffice", timeout=180)

    with patch("app.core.office_convert.subprocess.run", side_effect=fake_run):
        resp = client.post(
            f"/api/v1/jobs/{job.id}/documents/{tdoc.id}/upload",
            files={"file": ("invoice.docx", b"not a real docx",
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 422, resp.text
    assert "took too long" in resp.json()["detail"]
