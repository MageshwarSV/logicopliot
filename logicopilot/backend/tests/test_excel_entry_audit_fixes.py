"""Three bugs found by an independent code audit of app/core/excel_entry.py:

1. A merged cell in a customer's own uploaded template (source="template") silently
   swallowed whatever value got written into it - delete_rows clears cell CONTENT but
   leaves a stale merged_cells range declared over the cleared area, and both Excel and
   openpyxl read any non-anchor cell inside a merge as blank regardless of what the raw
   XML says is there. Contradicts the module's own "EVERYTHING HERE FAILS LOUDLY" promise.

2. Renaming a mark or custom field never updated excel_config's own "field": "<label>"
   mappings, silently blanking that column on every job's export forever -
   edit_custom_field's own docstring claims "Renaming is safe", which this made false.

3. A line-scoped sheet (scope="line") always emitted at least one row, even for a job with
   genuinely zero line items, because _rows_for_plan's row count started at the shared
   default of 1 and was never allowed to drop below it for this scope specifically.
"""
import tempfile
from pathlib import Path

import openpyxl

from app.core.excel_entry import build_for_job, rename_field_in_excel_config, _rows_for_plan


# ---- merged cells --------------------------------------------------------------------------

def _make_template_with_merge(tmp_path: Path) -> Path:
    """A' 'Consignee' title cell merged A1:A3, spanning the header row down into the data
    rows - an ordinary real-world pattern (a banner/title cell over a whole column)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GENERAL"
    ws["A1"] = "Consignee"
    ws.merge_cells("A1:A3")
    ws["B1"] = "HSCode"
    path = tmp_path / "template.xlsx"
    wb.save(str(path))
    return path


def test_a_value_written_into_a_merged_cell_is_not_silently_lost(tmp_path):
    template_path = _make_template_with_merge(tmp_path)
    config = {
        "source": "template",
        "sheets": [{
            "sheet": "GENERAL", "header_row": 1, "scope": "job",
            "columns": [
                {"column": "A", "header": "Consignee", "field": "consignee"},
                {"column": "B", "header": "HSCode", "field": "hs_code"},
            ],
        }],
    }
    out = build_for_job(
        config, {"consignee": "ACME IMPORTS PVT LTD", "hs_code": "854231"}, {},
        out_dir=tmp_path, template_path=template_path,
    )

    reopened = openpyxl.load_workbook(str(out))
    ws = reopened["GENERAL"]
    assert ws["A2"].value == "ACME IMPORTS PVT LTD"
    assert ws["B2"].value == "854231"


# ---- rename_field_in_excel_config ------------------------------------------------------------

def test_rename_rewrites_a_single_sheet_mapping():
    config = {"sheet": "GENERAL", "header_row": 1,
             "columns": [{"column": "A", "field": "old_label"},
                         {"column": "B", "field": "unrelated"}]}
    changed = rename_field_in_excel_config(config, "old_label", "new_label")
    assert changed is True
    assert config["columns"][0]["field"] == "new_label"
    assert config["columns"][1]["field"] == "unrelated"  # untouched


def test_rename_rewrites_every_matching_column_across_multiple_sheets():
    config = {"sheets": [
        {"sheet": "GENERAL", "columns": [{"column": "A", "field": "hs_code"}]},
        {"sheet": "ITEMS", "columns": [{"column": "C", "field": "hs_code"},
                                       {"column": "D", "field": "qty"}]},
    ]}
    changed = rename_field_in_excel_config(config, "hs_code", "cth_code")
    assert changed is True
    assert config["sheets"][0]["columns"][0]["field"] == "cth_code"
    assert config["sheets"][1]["columns"][0]["field"] == "cth_code"
    assert config["sheets"][1]["columns"][1]["field"] == "qty"


def test_rename_is_a_no_op_when_nothing_matches():
    config = {"sheet": "GENERAL", "columns": [{"column": "A", "field": "unrelated"}]}
    changed = rename_field_in_excel_config(config, "old_label", "new_label")
    assert changed is False
    assert config["columns"][0]["field"] == "unrelated"


def test_rename_handles_a_blank_or_missing_config():
    assert rename_field_in_excel_config(None, "a", "b") is False
    assert rename_field_in_excel_config({}, "a", "b") is False


# ---- _rows_for_plan: no phantom row on a genuinely empty line-scoped sheet -------------------

def test_a_line_scoped_sheet_with_no_line_items_produces_zero_rows():
    plan = {"scope": "line", "header_row": 1,
           "columns": [{"column": "A", "field": "item_description"},
                       {"column": "B", "field": "item_qty"}]}
    data = _rows_for_plan(plan, values={}, rows={})
    assert data == []


def test_a_line_scoped_sheet_with_real_line_items_is_unaffected():
    plan = {"scope": "line", "header_row": 1,
           "columns": [{"column": "A", "field": "item_description"},
                       {"column": "B", "field": "item_qty"}]}
    rows = {"item_description": ["Widget A", "Widget B"], "item_qty": ["10", "20"]}
    data = _rows_for_plan(plan, values={}, rows=rows)
    assert len(data) == 2


def test_a_job_scoped_sheet_still_always_produces_exactly_one_row():
    """Unaffected by the line-scope fix - scope="job" never depended on the shared default
    the same way, and must keep producing its one row regardless of line-item data."""
    plan = {"scope": "job", "header_row": 1,
           "columns": [{"column": "A", "field": "consignee"}]}
    data = _rows_for_plan(plan, values={"consignee": "ACME"}, rows={})
    assert len(data) == 1


# ---- _rows_for_plan: a column's own "transform", e.g. ContainerSize deriving its value ------- #
# from container_type ("40HC" -> "40") without needing its own separate CustomField/FieldMark.

def test_leading_digits_transform_splits_the_size_off_a_container_type_code():
    plan = {"scope": "line", "header_row": 1,
           "columns": [
               {"column": "A", "field": "container_type"},
               {"column": "B", "field": "container_type", "transform": "leading_digits"},
           ]}
    rows = {"container_type": ["40HC", "20GP", "45RF"]}
    data = _rows_for_plan(plan, values={}, rows=rows)
    # Column A (no transform) is unaffected - the raw value, exactly as every existing
    # mapping already behaves.
    assert [r[0] for r in data] == ["40HC", "20GP", "45RF"]
    # Column B (transformed) carries only the leading numeric size.
    assert [r[1] for r in data] == ["40", "20", "45"]


def test_leading_digits_transform_is_blank_when_nothing_at_the_front_is_numeric():
    plan = {"scope": "line", "header_row": 1,
           "columns": [{"column": "A", "field": "container_type", "transform": "leading_digits"}]}
    rows = {"container_type": ["", "GP"]}
    data = _rows_for_plan(plan, values={}, rows=rows)
    assert [r[0] for r in data] == ["", ""]


def test_an_unknown_transform_name_is_ignored_not_raised():
    plan = {"scope": "job", "header_row": 1,
           "columns": [{"column": "A", "field": "consignee", "transform": "not_a_real_transform"}]}
    data = _rows_for_plan(plan, values={"consignee": "ACME"}, rows={})
    assert data == [["ACME"]]


def test_transform_applies_to_a_job_level_value_too_not_only_per_line_rows():
    plan = {"scope": "job", "header_row": 1,
           "columns": [{"column": "A", "field": "container_type", "transform": "leading_digits"}]}
    data = _rows_for_plan(plan, values={"container_type": "40HC"}, rows={})
    assert data == [["40"]]
